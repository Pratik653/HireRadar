"""Excel report: Hiring Leads, Search Summary and Application Tracker sheets."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.models import ExportRequest, Opportunity, TrackedStatus
from app.scoring import TFIDF_SATURATION
from app.security import maybe_safe_url

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="1D4ED8", underline="single")
MISSING = "—"
STATUSES = ["not_applied", "saved", "applied", "contacted", "interviewing", "offer", "rejected"]


def _put(ws: Worksheet, row: int, col: int, value, *, wrap: bool = False):
    """Write a cell. Text that a spreadsheet would read as a formula is stored
    as a literal string, so imported post text can never execute."""
    if value is None or value == "":
        value = MISSING
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        cell.data_type = "s"
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    else:
        cell.alignment = Alignment(vertical="top")
    return cell


def _link(ws: Worksheet, row: int, col: int, url: str | None, label: str):
    safe = maybe_safe_url(url)
    if not safe:
        return _put(ws, row, col, None)
    cell = _put(ws, row, col, label)
    cell.hyperlink = safe
    cell.font = LINK_FONT
    return cell


def _header(ws: Worksheet, headers: list[tuple[str, int]]) -> None:
    for col, (name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _finish_table(ws: Worksheet, columns: int, rows: int) -> None:
    ws.auto_filter.ref = f"A1:{get_column_letter(columns)}{max(rows, 1) + 1}"


def _date(opp: Opportunity) -> str | None:
    return opp.published_at.strftime("%Y-%m-%d %H:%M UTC") if opp.published_at else None


def _leads(ws: Worksheet, results: list[Opportunity], tracking: dict[str, TrackedStatus]) -> None:
    headers = [
        ("Score", 8),
        ("Role", 34),
        ("Company", 24),
        ("Contact", 22),
        ("Contact position", 26),
        ("Contact source", 14),
        ("Country", 18),
        ("State / Province", 18),
        ("City", 14),
        ("Location as listed", 26),
        ("Work mode", 11),
        ("Published", 19),
        ("Date confidence", 15),
        ("Hiring intent", 30),
        ("Source", 20),
        ("Post URL", 12),
        ("Contact profile", 14),
        ("Matched skills", 26),
        ("Evidence", 60),
        ("Why it ranked", 70),
        ("Duplicates merged", 10),
        ("Application status", 16),
        ("Data", 10),
    ]
    _header(ws, headers)
    for row, opp in enumerate(results, start=2):
        contact = opp.contact
        values = [
            opp.score,
            opp.title,
            opp.company,
            contact.name if contact else None,
            contact.position if contact else None,
            contact.provenance.replace("_", " ") if contact else None,
            opp.country,
            opp.subdivision,
            opp.city,
            opp.location_raw,
            opp.work_mode,
            _date(opp),
            opp.date_confidence,
            opp.intent_label,
            opp.source_label,
        ]
        for col, value in enumerate(values, start=1):
            _put(ws, row, col, value)
        _link(ws, row, 16, opp.source_url, "Open post")
        _link(ws, row, 17, contact.profile_url if contact else None, "Profile")
        _put(ws, row, 18, ", ".join(opp.matched_skills))
        _put(ws, row, 19, opp.snippet, wrap=True)
        _put(ws, row, 20, "\n".join(opp.explanations), wrap=True)
        _put(ws, row, 21, len(opp.duplicates))
        _put(ws, row, 22, tracking.get(opp.id, TrackedStatus()).status.replace("_", " "))
        _put(ws, row, 23, "SYNTHETIC" if opp.is_demo else "live")
    _finish_table(ws, len(headers), len(results))


def _summary(ws: Worksheet, payload: ExportRequest) -> None:
    response = payload.response
    req = response.request
    s = response.summary
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    locations = "; ".join(loc.label for loc in req.locations)
    rows: list[tuple[str, object]] = [
        ("Mode", "DEMO — synthetic data, not real opportunities" if response.mode == "demo" else "LIVE"),
        ("Generated", response.generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Locations", locations or "any"),
        ("Work modes", ", ".join(req.work_modes)),
        ("Job titles", ", ".join(req.titles)),
        ("Skills", ", ".join(req.skills)),
        ("Exclusions", ", ".join(req.exclusions)),
        ("Seniority", ", ".join(req.seniority) or "any"),
        ("Publication window", f"last {req.days} day(s)"),
        ("Undated items", "included (recency score 0)" if req.include_undated else "excluded"),
        ("Minimum score", req.min_score),
        ("", ""),
        ("Items fetched", s.fetched),
        ("Results", s.returned),
        ("Removed: duplicates", s.removed_duplicates),
        ("Removed: exclusion keywords", s.removed_by_exclusion),
        ("Removed: location / work mode", s.removed_by_location),
        ("Removed: outside date window", s.removed_by_date),
        ("Removed: not hiring content", s.removed_unrelated),
        ("Removed: below minimum score", s.removed_below_min_score),
        ("Removed: role relevance under 15%", s.removed_low_relevance),
        ("Results without a publication date", s.undated),
        ("", ""),
    ]
    for status in response.sources:
        detail = f"{status.status}, {status.fetched} fetched" + (f" — {status.message}" if status.message else "")
        rows.append((f"Source: {status.label}", detail))
    rows += [
        ("", ""),
        (
            "Scoring",
            "score = 100 × (role × {role} + intent × {intent} + geography × {geography} + recency × {recency})".format(
                **response.weights
            ),
        ),
        (
            "Role",
            "Title-word overlap, share of your skills found (with synonyms), and TF-IDF cosine "
            f"similarity to your titles/skills/description (cosine ≥ {TFIDF_SATURATION} counts as full).",
        ),
        (
            "Intent",
            "Rule-based: explicit vacancy 1.0, hiring manager 0.95, recruiter 0.85, general 0.7, weak signal 0.3.",
        ),
        (
            "Geography",
            "Inside a selected region, country, state or city 1.0; same state, other city 0.7; same country, "
            "other city or state 0.6; remote 0.5–1.0 by stated eligibility; unstated 0.3. Selections are OR-ed.",
        ),
        ("Recency", "Linear from 1.0 (just published) to 0 (end of window). Unknown dates score 0."),
        ("Caveat", "Weights are heuristics, not validated accuracy measures. Verify every lead at its source."),
    ]
    for note in response.notes:
        rows.append(("Note", note))
    for index, (key, value) in enumerate(rows, start=1):
        if not key:
            continue
        _put(ws, index, 1, key).font = Font(bold=True)
        _put(ws, index, 2, value, wrap=True)


def _tracker(ws: Worksheet, results: list[Opportunity], tracking: dict[str, TrackedStatus]) -> None:
    headers = [
        ("Opportunity", 36),
        ("Company", 24),
        ("Contact", 22),
        ("Post URL", 12),
        ("Status", 16),
        ("Follow-up date", 15),
        ("Notes", 50),
        ("Score", 8),
    ]
    _header(ws, headers)
    validation = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(validation)
    for row, opp in enumerate(results, start=2):
        tracked = tracking.get(opp.id, TrackedStatus())
        _put(ws, row, 1, opp.title)
        _put(ws, row, 2, opp.company)
        _put(ws, row, 3, opp.contact.name if opp.contact else None)
        _link(ws, row, 4, opp.source_url, "Open post")
        status_cell = ws.cell(row=row, column=5, value=tracked.status)
        validation.add(status_cell)
        _put(ws, row, 6, tracked.follow_up)
        _put(ws, row, 7, tracked.notes, wrap=True)
        _put(ws, row, 8, opp.score)
    _finish_table(ws, len(headers), len(results))


def build_workbook(payload: ExportRequest) -> bytes:
    wb = Workbook()
    leads = wb.active
    leads.title = "Hiring Leads"
    results = payload.response.results
    _leads(leads, results, payload.tracking)
    _summary(wb.create_sheet("Search Summary"), payload)
    _tracker(wb.create_sheet("Application Tracker"), results, payload.tracking)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
