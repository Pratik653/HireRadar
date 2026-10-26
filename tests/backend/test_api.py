"""API contracts, validation, export and security behaviour. No network access:
live adapters are exercised with a stubbed HTTP transport."""

from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.models import ImportedPost, SearchRequest
from app.search import run_search
from app.security import RateLimiter, safe_public_url
from app.sources.http import SafeClient, SourceError

DEMO = {
    "mode": "demo",
    "titles": ["Applied AI Engineer", "LLM Engineer"],
    "skills": ["Python", "RAG", "LangGraph"],
    "locations": [{"country_code": "AE"}, {"country_code": "QA", "city": "Doha"}],
    "days": 30,
}


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def test_health(client):
    body = client.get("/health").json()
    assert body["status"] == "ok" and body["version"]


def test_meta_lists_locations_and_sources(client):
    body = client.get("/api/meta").json()
    locations = body["locations"]
    assert locations["countries"] > 200 and locations["subdivisions"] > 100 and locations["cities"] > 250
    assert {"world", "asia", "europe", "gcc"} <= {r["id"] for r in locations["regions"]}
    assert {s["key"] for s in body["sources"]} == {"manual_import", "company_boards", "remotive"}
    assert body["demo_source"]["live"] is False


def test_demo_search_is_synthetic_and_filtered(client):
    body = client.post("/api/search", json=DEMO).json()
    assert body["mode"] == "demo"
    assert body["results"], "demo should return results"
    assert all(r["is_demo"] for r in body["results"])
    assert all(r["source_url"].startswith("https://example.org/") for r in body["results"])
    assert {r["country"] for r in body["results"]} <= {"United Arab Emirates", "Qatar", None}
    assert body["summary"]["removed_unrelated"] >= 2  # job-seeker and announcement posts
    assert body["summary"]["removed_duplicates"] >= 1
    scores = [r["score"] for r in body["results"]]
    assert scores == sorted(scores, reverse=True)


def test_live_search_never_returns_demo_data(client):
    payload = {
        "mode": "live",
        "titles": ["Python Developer"],
        "sources": ["manual_import"],
        "imports": [
            {
                "text": "We're hiring a Python Developer in Dubai. Apply now via our site.",
                "url": "https://www.linkedin.com/posts/example-123",
                "author_name": "A. Person",
                "published_at": "2026-10-25T09:00:00Z",
            }
        ],
    }
    body = client.post("/api/search", json=payload).json()
    assert body["mode"] == "live"
    assert not any(r["is_demo"] for r in body["results"])
    assert body["results"][0]["contact"]["provenance"] == "user_provided"
    assert body["results"][0]["date_confidence"] == "reported"


def test_validation_requires_a_target(client):
    response = client.post("/api/search", json={"mode": "demo"})
    assert response.status_code == 422
    assert response.json()["detail"] == "Invalid request"


@pytest.mark.parametrize(
    "payload",
    [
        {**DEMO, "days": 5},
        {**DEMO, "locations": [{"country_code": "XX"}]},
        {**DEMO, "locations": [{"country_code": "AE", "city": "Doha"}]},
        {**DEMO, "company_boards": [{"provider": "greenhouse", "slug": "../../etc"}]},
        {**DEMO, "unexpected": True},
    ],
)
def test_validation_rejects_bad_input(client, payload):
    assert client.post("/api/search", json=payload).status_code == 422


def test_body_size_limit(client):
    big = {**DEMO, "description": "x" * 300_000}
    assert client.post("/api/search", json=big).status_code == 413


def test_cors_allows_only_configured_origins(client):
    ok = client.options(
        "/api/search", headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "POST"}
    )
    bad = client.options(
        "/api/search", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST"}
    )
    assert ok.headers.get("access-control-allow-origin") == "http://localhost:5173"
    assert "access-control-allow-origin" not in bad.headers


def test_rate_limiter_window():
    limiter = RateLimiter(per_minute=2)
    assert limiter.allow("ip", now=0) and limiter.allow("ip", now=1)
    assert not limiter.allow("ip", now=2)
    assert limiter.allow("ip", now=61.5)


@pytest.mark.parametrize(
    "url",
    [
        "javascript:alert(1)",
        "data:text/html,hi",
        "file:///etc/passwd",
        "http://127.0.0.1/admin",
        "http://169.254.169.254/latest",
        "https://user:pw@example.com",
        "http://localhost:8000",
        "http://intranet",
        "http://printer.local/",
    ],
)
def test_unsafe_urls_rejected(url):
    with pytest.raises(ValueError):
        safe_public_url(url)


def test_safe_url_is_normalised():
    assert safe_public_url("HTTPS://www.LinkedIn.com/posts/abc#frag") == "https://www.linkedin.com/posts/abc"


def test_imported_post_rejects_unsafe_link():
    with pytest.raises(ValueError):
        ImportedPost(text="We are hiring a Python developer in Dubai", url="javascript:alert(1)")


async def test_outbound_client_enforces_allowlist():
    client = SafeClient()
    with pytest.raises(SourceError, match="allowlist"):
        await client.get_json("https://evil.example.com/jobs")
    with pytest.raises(SourceError, match="allowlist"):
        await client.get_json("http://boards-api.greenhouse.io/v1/boards/x/jobs")
    await client.aclose()


def _stub(handler) -> SafeClient:
    return SafeClient(httpx.AsyncClient(transport=httpx.MockTransport(handler)))


async def test_company_board_adapter_normalises_greenhouse():
    def handler(req: httpx.Request) -> httpx.Response:
        assert req.url.host == "boards-api.greenhouse.io"
        return httpx.Response(
            200,
            json={
                "jobs": [
                    {
                        "title": "Senior Python Engineer",
                        "company_name": "Careem",
                        "absolute_url": "https://boards.greenhouse.io/careem/jobs/1",
                        "location": {"name": "Dubai, United Arab Emirates"},
                        "first_published": "2026-10-24T08:00:00-04:00",
                        "content": "&lt;p&gt;Python, FastAPI and RAG systems&lt;/p&gt;",
                    }
                ]
            },
        )

    req = SearchRequest(
        titles=["Python Engineer"],
        skills=["Python"],
        sources=["company_boards"],
        company_boards=[{"provider": "greenhouse", "slug": "careem"}],
        days=30,
    )
    response = await run_search(req, _stub(handler))
    status = response.sources[0]
    assert status.status == "ok" and status.fetched == 1 and status.live
    job = response.results[0]
    assert job.company == "Careem" and job.city == "Dubai" and job.date_confidence == "exact"
    assert job.intent.value == "explicit_vacancy"
    assert "Python, FastAPI" in job.snippet


async def test_source_failure_is_reported_not_replaced_with_demo():
    def handler(req: httpx.Request) -> httpx.Response:
        return httpx.Response(503)

    req = SearchRequest(titles=["Python"], sources=["remotive"], work_modes=["remote"])
    response = await run_search(req, _stub(handler))
    assert response.sources[0].status == "error"
    assert "503" in response.sources[0].message
    assert response.results == []


async def test_missing_board_reports_not_found():
    req = SearchRequest(
        titles=["Python"], sources=["company_boards"], company_boards=[{"provider": "lever", "slug": "nope"}]
    )
    response = await run_search(req, _stub(lambda r: httpx.Response(404)))
    assert response.sources[0].status == "error" and "not found" in response.sources[0].message


def test_export_workbook(client):
    search = client.post("/api/search", json=DEMO).json()
    first = search["results"][0]
    first["title"] = '=HYPERLINK("http://evil.example","click")'
    first["snippet"] = "+cmd|' /C calc'!A0"
    tracking = {first["id"]: {"status": "applied", "follow_up": "2026-11-01", "notes": "@SUM(1)"}}
    response = client.post("/api/export", json={"response": search, "tracking": tracking})
    assert response.status_code == 200
    assert "hireradar-demo-" in response.headers["content-disposition"]

    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Hiring Leads", "Search Summary", "Application Tracker"]
    leads = wb["Hiring Leads"]
    assert leads.freeze_panes == "A2" and leads.auto_filter.ref.startswith("A1:")
    headers = [c.value for c in leads[1]]
    assert headers[:3] == ["Score", "Role", "Company"]

    title_cell = leads.cell(row=2, column=2)
    assert title_cell.data_type == "s" and title_cell.value.startswith("=HYPERLINK")
    assert leads.cell(row=2, column=19).data_type == "s"
    assert leads.cell(row=2, column=16).hyperlink.target.startswith("https://example.org/")
    assert leads.cell(row=2, column=23).value == "SYNTHETIC"

    tracker = wb["Application Tracker"]
    assert tracker.cell(row=2, column=5).value == "applied"
    assert tracker.cell(row=2, column=7).data_type == "s"
    summary = {r[0].value: r[1].value for r in wb["Search Summary"].iter_rows() if r[0].value}
    assert summary["Mode"].startswith("DEMO")
    assert "role × 0.4" in summary["Scoring"]


def test_export_missing_values_use_placeholder(client):
    search = client.post("/api/search", json=DEMO).json()
    first = search["results"][0]
    first["contact"] = None
    first["source_url"] = "javascript:alert(1)"
    wb = load_workbook(BytesIO(client.post("/api/export", json={"response": search}).content))
    leads = wb["Hiring Leads"]
    assert leads.cell(row=2, column=4).value == "—"
    assert leads.cell(row=2, column=16).hyperlink is None
